"""Build DESIGN.docx from DESIGN.md.

Headings, bullet lists, tables and code blocks are mapped onto python-docx
styles so the result is a real Word document rather than styled plain text.
East-Asian font is set explicitly: font.name alone does not cover CJK glyphs.
"""
import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

SRC = Path(sys.argv[1])
OUT = Path(sys.argv[2])
CJK = "Microsoft YaHei"

# A4 with 2cm margins. Without an explicit width the tables below stretch to fit
# their longest cell and push the whole text column off the page.
PAGE_W_CM = 21.0
MARGIN_CM = 2.0
CONTENT_CM = PAGE_W_CM - 2 * MARGIN_CM


def style_run(run, *, bold=None, size=None, mono=False, color=None):
    run.font.name = "Consolas" if mono else "Calibri"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), CJK)
    if bold is not None:
        run.bold = bold
    if size is not None:
        run.font.size = Pt(size)
    if color is not None:
        run.font.color.rgb = RGBColor(*color)


def add_inline(paragraph, text, *, mono=False):
    """Turn **bold** and `code` into runs, leaving the rest as plain text."""
    for part in re.split(r"(\*\*[^*]+\*\*|`[^`]+`)", text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            style_run(paragraph.add_run(part[2:-2]), bold=True, mono=mono)
        elif part.startswith("`") and part.endswith("`"):
            style_run(paragraph.add_run(part[1:-1]), mono=True)
        else:
            style_run(paragraph.add_run(part), mono=mono)


def flush_table(doc, rows):
    if not rows:
        return
    header, body = rows[0], rows[1:]
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Light Grid Accent 1"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        cell.text = ""
        add_inline(cell.paragraphs[0], text)
        for run in cell.paragraphs[0].runs:
            run.bold = True
    for row in body:
        cells = table.add_row().cells
        for cell, text in zip(cells, row):
            cell.text = ""
            add_inline(cell.paragraphs[0], text)
    _fix_table_width(table, len(header))
    doc.add_paragraph()


def _fix_table_width(table, columns):
    """Pin the table to the text column.

    python-docx leaves autofit on, and LibreOffice then sizes the table to its
    longest cell -- wider than the page -- which drags the whole text column out
    of the margins. `autofit = False` writes `w:tblLayout` in the schema-correct
    position; the per-cell widths do the rest. Do not append `w:tblW` by hand:
    `w:tblPr` children must appear in schema order, and appending puts it after
    `w:jc`.
    """
    table.autofit = False

    # A narrow first column reads as a label column, which suits every table in
    # this document ("优先级 | 功能", "能力 | 方案 | ...").
    if columns == 1:
        fractions = [1.0]
    else:
        first = 0.20 if columns >= 3 else 0.22
        rest = (1.0 - first) / (columns - 1)
        fractions = [first] + [rest] * (columns - 1)

    for row in table.rows:
        for cell, fraction in zip(row.cells, fractions):
            cell.width = Cm(CONTENT_CM * fraction)

    # `add_table` writes a `tblGrid` of equal columns and nothing updates it, so
    # the grid and the per-cell widths would disagree. Renderers differ on which
    # one they honour, so make them say the same thing.
    grid = table._tbl.find(qn("w:tblGrid"))
    if grid is not None:
        cols = grid.findall(qn("w:gridCol"))
        for col, fraction in zip(cols, fractions):
            col.set(qn("w:w"), str(int(CONTENT_CM / 2.54 * 1440 * fraction)))


def build():
    doc = Document()
    normal = doc.styles["Normal"]
    normal.font.size = Pt(10.5)
    normal.font.name = "Calibri"
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), CJK)

    for section in doc.sections:
        section.orientation = WD_ORIENT.PORTRAIT
        section.page_width = Cm(PAGE_W_CM)
        section.page_height = Cm(29.7)
        section.left_margin = Cm(MARGIN_CM)
        section.right_margin = Cm(MARGIN_CM)
        section.top_margin = Cm(MARGIN_CM)
        section.bottom_margin = Cm(MARGIN_CM)

    lines = SRC.read_text(encoding="utf-8").splitlines()
    i = 0
    table_buf = []
    code_buf = []
    in_code = False

    while i < len(lines):
        line = lines[i]

        # fenced code block
        if line.strip().startswith("```"):
            if in_code:
                para = doc.add_paragraph()
                para.paragraph_format.left_indent = Pt(12)
                para.paragraph_format.space_after = Pt(6)
                run = para.add_run("\n".join(code_buf))
                style_run(run, mono=True, size=9)
                code_buf = []
                in_code = False
            else:
                flush_table(doc, table_buf)
                table_buf = []
                in_code = True
            i += 1
            continue
        if in_code:
            code_buf.append(line)
            i += 1
            continue

        # table rows
        if line.strip().startswith("|") and line.strip().endswith("|"):
            cells = [c.strip() for c in line.strip().strip("|").split("|")]
            if all(re.fullmatch(r":?-{2,}:?", c) for c in cells):
                i += 1
                continue
            table_buf.append(cells)
            i += 1
            continue
        flush_table(doc, table_buf)
        table_buf = []

        stripped = line.strip()

        if not stripped:
            i += 1
            continue

        if stripped == "---":
            doc.add_paragraph("─" * 40).alignment = WD_ALIGN_PARAGRAPH.CENTER
            i += 1
            continue

        # headings
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            level = len(m.group(1))
            text = m.group(2)
            if level == 1:
                doc.add_heading(text, level=0)
            else:
                doc.add_heading(text, level=min(level - 1, 4))
            i += 1
            continue

        # blockquote
        if stripped.startswith(">"):
            para = doc.add_paragraph(style="Intense Quote")
            add_inline(para, stripped.lstrip("> ").strip())
            i += 1
            continue

        # list items
        m = re.match(r"^(\s*)([-*]|\d+\.)\s+(.*)$", line)
        if m:
            indent = len(m.group(1)) // 2
            text = m.group(3)
            style = "List Number" if m.group(2)[0].isdigit() else "List Bullet"
            try:
                para = doc.add_paragraph(style=style if indent == 0 else f"{style} {min(indent + 1, 3)}")
            except KeyError:
                para = doc.add_paragraph(style=style)
            add_inline(para, text)
            i += 1
            continue

        para = doc.add_paragraph()
        add_inline(para, stripped)
        i += 1

    flush_table(doc, table_buf)
    doc.save(OUT)
    print(f"written: {OUT}")


build()
